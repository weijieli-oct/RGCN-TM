
import csv
import json
import math
from pathlib import Path
from datetime import datetime
import numpy as np
import params as settings

TRAIN_RESULT_DIRECTORY=None
PLOT_SIDES=None
PLOT_TIMES=[10.,30.,50.,100.]
PLOT_DISCARD_TIME=50.
FIGURE_DIRECTORY=None
VORTICITY_START=1.
VORTICITY_END=100.
VORTICITY_INTERVAL=1.
VORTICITY_FPS=10
VORTICITY_LIMIT=None
VORTICITY_DPI=100
MAKE_VORTICITY_GIF=True
PHASE_SEARCH_START=50.
PHASE_SEARCH_END=None
PHASE_MANUAL_CYCLES={}
PHASE_MAX_TIME_ERROR_FRACTION=.10
PHASE_MAX_PERIOD_VARIATION=.15
PHASE_MIN_LIFT_STD=1e-6
AMPLITUDE_START=70.
AMPLITUDE_END=100.


def shape_name(n):
    return "circle" if int(n)==0 else f"polygon_{int(n)}"


def requested_times(start,end,interval):
    if not all(math.isfinite(v) for v in(start,end,interval)) or start<0 or end<start or interval<=0:
        raise ValueError("Invalid output time range")
    steps=(end-start)/interval
    if not math.isclose(steps,round(steps),abs_tol=1e-6,rel_tol=0): raise ValueError("Output range must align with interval")
    return [float(start+i*interval) for i in range(round(steps)+1)]


def read_rows(path):
    if not path.exists(): return []
    with path.open(encoding="utf-8-sig") as f: return list(csv.DictReader(f))


def pyplot():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


def draw_body(ax,n):
    plt=pyplot()
    if int(n)==0: ax.add_patch(plt.Circle((0,0),.5,color="gray")); return
    n=int(n); radius=math.sqrt(math.pi/(2*n*math.sin(2*math.pi/n)))
    a=math.pi/2-math.pi/n+2*math.pi*np.arange(n)/n
    ax.add_patch(plt.Polygon(radius*np.stack((np.cos(a),np.sin(a)),axis=1),color="gray"))


def read_field(path,case,n):
    with np.load(path,allow_pickle=False) as data:
        if "sides" not in data or int(data["sides"][case])!=n: raise ValueError("Wrong geometry snapshot")
        x,y=data["x"].copy(),data["y"].copy(); w=data["vorticity"][case,0].astype(float)
        solid=data["solid_mask"][case].astype(bool)
        if not np.isfinite(w[~solid]).all(): raise ValueError("Nonfinite fluid vorticity")
        w[solid]=np.nan
        return x,y,w,float(data["time"])


def inventory(root,case,n):
    fields=[]
    for path in root.glob("field_*.npz"):
        with np.load(path,allow_pickle=False) as data:
            if "sides" not in data: raise ValueError("Results from the legacy cylindrical Re version cannot be read as geometry-version results")
            if case<len(data["sides"]) and int(data["sides"][case])==n: fields.append((float(data["time"]),path))
    fields.sort()
    if len({t for t,p in fields})!=len(fields): raise ValueError("Duplicate snapshot times")
    return fields


def plot_vorticity_sequence(fields,case,n,destination,times,fps=10,make_gif=True,limit=None,dpi=100,interval=1.):
    from PIL import Image
    if not math.isfinite(fps) or not 0<fps<=100 or dpi<=0: raise ValueError("Invalid animation FPS/DPI")
    if limit is not None and (not math.isfinite(limit) or limit<=0): raise ValueError("Positive color limit required")
    selected=[]; missing=[]
    for target in times:
        found=next(((t,p) for t,p in fields if math.isclose(t,target,abs_tol=1e-6,rel_tol=0)),None)
        if found is None: missing.append(target)
        else: selected.append(found)
    report={"sides":n,"requested_times":times,"missing_times":missing,"saved_times":[t for t,p in selected],
            "note":"Actual saved fields only; no temporal interpolation"}
    destination.mkdir(parents=True,exist_ok=True)
    if not selected:
        (destination/"sequence_report.json").write_text(json.dumps(report,indent=2),encoding="utf-8"); return report
    clim=float(limit) if limit is not None else max(1e-5,max(float(np.nanmax(np.abs(read_field(p,case,n)[2]))) for t,p in selected))
    report["color_limit"]=clim; plt=pyplot(); frames=[]; durations=[]
    try:
        for i,(t,p) in enumerate(selected):
            x,y,w,t=read_field(p,case,n)
            fig,ax=plt.subplots(figsize=(10,4))
            mesh=ax.pcolormesh(x,y,w,cmap="RdBu_r",shading="auto",vmin=-clim,vmax=clim)
            draw_body(ax,n); ax.set_aspect("equal"); ax.set_xlabel("x/D_eq"); ax.set_ylabel("y/D_eq")
            ax.set_title(f"{shape_name(n)}, Re=100, t*={t:.4f}")
            fig.colorbar(mesh,ax=ax,label="vorticity D_eq/U_ref"); fig.tight_layout()
            png=destination/f"vorticity_{i:05d}_t_{t:.4f}.png"; fig.savefig(png,dpi=dpi); plt.close(fig)
            if make_gif:
                with Image.open(png) as im:
                    frames.append(im.convert("RGB").convert("P",palette=getattr(Image,"Palette",Image).ADAPTIVE))
                gap=selected[i+1][0]-t if i+1<len(selected) else interval
                durations.append(max(10,int(round(100*gap/interval/fps))*10))
            if(i+1)%20==0: print(shape_name(n),"vorticity frames",i+1,"/",len(selected),flush=True)
        if len(frames)>=2:
            name="vorticity_partial.gif" if missing else "vorticity_animation.gif"
            frames[0].save(destination/name,save_all=True,append_images=frames[1:],duration=durations,loop=0,optimize=False,disposal=2)
            report["animation"]=name
    finally:
        for frame in frames: frame.close()
    (destination/"sequence_report.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
    return report


def force_plots(root,case,n,destination,discard_time):
    plt=pyplot(); rows=[r for r in read_rows(root/"force_history.csv") if int(r["case"])==case]
    if not rows: return
    rows.sort(key=lambda r:float(r["time"])); t=np.array([float(r["time"]) for r in rows])
    if np.any(np.diff(t)<=0): raise ValueError("Non-increasing force history")
    array=lambda key:np.array([float(r[key]) for r in rows])
    for filename,suffixes in (("force_coefficients",("total","pressure","viscous")),
                              ("force_method_comparison",("total","shifted_estimate","penalty_trapezoid","penalty_inertia_indicator"))):
        fig,axes=plt.subplots(2,1,figsize=(10,7),sharex=True)
        for ax,prefix in zip(axes,("Cd","Cl")):
            for suffix in suffixes: ax.plot(t,array(prefix+"_"+suffix),label=suffix)
            ax.set_ylabel(prefix); ax.legend(fontsize=8)
        axes[-1].set_xlabel("t*=t U_ref/D_eq"); fig.suptitle(shape_name(n)+", Re=100; reconstructed estimates")
        fig.tight_layout(); fig.savefig(destination/(filename+".png"),dpi=160); plt.close(fig)
    history=[r for r in read_rows(root/"history.csv") if int(r["case"])==case]
    fig,axes=plt.subplots(2,1,figsize=(10,6),sharex=True)
    for ax,prefix in zip(axes,("Cd","Cl")):
        ax.plot([float(r["time"]) for r in history],[float(r[prefix+"_penalty_indicator"]) for r in history]); ax.set_ylabel(prefix+" penalty indicator")
    axes[-1].set_xlabel("t*"); fig.tight_layout(); fig.savefig(destination/"force_indicators.png",dpi=160); plt.close(fig)
    mask=t>=discard_time; times=t[mask]; report={"sides":n,"Re":100,"discard_time":discard_time,
        "note":"Window statistics and reconstruction sensitivity; no reference-error claim"}
    if len(times)>=2:
        trap=getattr(np,"trapezoid",None) or np.trapz
        for key in rows[0]:
            if not key.startswith(("Cd_","Cl_")): continue
            values=array(key)[mask]
            if not np.isfinite(values).all(): continue
            mean=float(trap(values,times)/(times[-1]-times[0]))
            report[key]={"mean":mean,"rms":float(np.sqrt(trap(values**2,times)/(times[-1]-times[0]))),
                         "fluctuation_rms":float(np.sqrt(trap((values-mean)**2,times)/(times[-1]-times[0]))),
                         "min":float(values.min()),"max":float(values.max())}
        cl=array("Cl_total")[mask]
        if len(times)>=32 and np.isfinite(cl).all() and np.std(cl)>1e-8 and np.allclose(np.diff(times),np.diff(times)[0],rtol=1e-4,atol=1e-7):
            detrended=cl-np.polyval(np.polyfit(times-times[0],cl,1),times-times[0])
            spectrum=np.abs(np.fft.rfft(detrended*np.hanning(len(cl))))**2
            f=np.fft.rfftfreq(len(cl),d=times[1]-times[0]); peak=1+int(np.argmax(spectrum[1:]))
            report["frequency"]={"candidate_St":float(f[peak]),"bin_width":float(f[1]),"candidate_cycles":float(f[peak]*(times[-1]-times[0]))}
            fig,ax=plt.subplots(figsize=(9,4)); ax.plot(f[1:],spectrum[1:]); ax.set_xlabel("St=f D_eq/U_ref"); ax.set_ylabel("Windowed power")
            fig.tight_layout(); fig.savefig(destination/"lift_stress_spectrum.png",dpi=160); plt.close(fig)
    (destination/"force_statistics.json").write_text(json.dumps(report,indent=2),encoding="utf-8")


def phase_plot(root,case,n,fields,destination):
    report={"sides":n,"Re":100,"source":"Cl_total","status":"not_available"}
    def done(reason):
        report["status"]=reason
        (destination/"cycle_report.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
        return report
    manual=PHASE_MANUAL_CYCLES.get(n)
    if manual is not None:
        if len(manual)!=2 or not all(math.isfinite(v) for v in manual) or manual[0]<0 or manual[1]<=0: raise ValueError("Invalid manual cycle")
        start,period=manual; report["method"]="user-specified"
    else:
        rows=[r for r in read_rows(root/"force_history.csv") if int(r["case"])==case and float(r["time"])>=PHASE_SEARCH_START
              and(PHASE_SEARCH_END is None or float(r["time"])<=PHASE_SEARCH_END)]
        rows.sort(key=lambda r:float(r["time"]))
        if len(rows)<32: return done("insufficient_lift_samples")
        t=np.array([float(r["time"]) for r in rows]); cl=np.array([float(r["Cl_total"]) for r in rows])
        if not np.isfinite(cl).all() or np.any(np.diff(t)<=0): return done("invalid_samples")
        if not np.allclose(np.diff(t),np.diff(t)[0],rtol=1e-4,atol=1e-7): return done("nonuniform_samples")
        if np.std(cl)<PHASE_MIN_LIFT_STD: return done("negligible_oscillation")
        cl=cl-cl.mean(); idx=np.flatnonzero((cl[:-1]<=0)&(cl[1:]>0))
        crosses=t[idx]-cl[idx]*np.diff(t)[idx]/(cl[idx+1]-cl[idx])
        if len(crosses)<4: return done("need_three_observed_cycles")
        periods=np.diff(crosses)[-3:]; variation=float(np.std(periods)/np.mean(periods))
        report["recent_periods"]=periods.tolist(); report["period_variation"]=variation
        if variation>PHASE_MAX_PERIOD_VARIATION: return done("irregular_periods")
        start,period=float(crosses[-2]),float(periods[-1]); report["method"]="upward mean-lift crossings"
        if period/np.diff(t)[0]<16: return done("undersampled_lift")
    report.update(cycle_start=start,period=period)
    if len(fields)<4: return done("insufficient_snapshots")
    chosen=[]
    for phase in(0.,.25,.5,.75):
        target=start+phase*period; actual,path=min(fields,key=lambda q:abs(q[0]-target))
        chosen.append({"target_phase":phase,"target_time":target,"actual_time":actual,"actual_phase":(actual-start)/period,
                       "phase_error":(actual-target)/period,"snapshot":str(path)})
    report["phases"]=chosen
    if len({r["snapshot"] for r in chosen})<4 or any(abs(r["phase_error"])>PHASE_MAX_TIME_ERROR_FRACTION for r in chosen): return done("snapshots_too_sparse")
    frames=[read_field(Path(r["snapshot"]),case,n) for r in chosen]
    limit=max(1e-5,max(float(np.nanmax(np.abs(w))) for x,y,w,t in frames))
    plt=pyplot(); fig,axes=plt.subplots(2,2,figsize=(15,8),constrained_layout=True)
    for ax,row,(x,y,w,t) in zip(axes.flat,chosen,frames):
        mesh=ax.pcolormesh(x,y,w,cmap="RdBu_r",vmin=-limit,vmax=limit,shading="auto"); draw_body(ax,n)
        ax.set_aspect("equal"); ax.set_xlabel("x/D_eq"); ax.set_ylabel("y/D_eq")
        ax.set_title(f"target {row['target_phase']:.2f}T; actual {row['actual_phase']:.3f}T; t*={t:.4f}")
    fig.suptitle(f"{shape_name(n)}, Re=100, candidate T*={period:.4f}")
    fig.colorbar(mesh,ax=list(axes.flat),label="vorticity D_eq/U_ref",shrink=.85)
    fig.savefig(destination/"four_phases.png",dpi=180); fig.savefig(destination/"four_phases.pdf"); plt.close(fig)
    return done("four_nearest_phase_snapshots")


def export_comparisons(roots,destination,discard_time=50.,sides=None):
    """Geometry is the grouping key everywhere; Re remains 100 in all exported rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    destination=settings.resolve_path(destination); destination.mkdir(parents=True,exist_ok=True)
    all_rows={name:[] for name in("history.csv","force_history.csv","probe_history.csv")}
    for root in roots:
        root=settings.resolve_path(root)
        for name in all_rows:
            for row in read_rows(root/name):
                if sides is None or int(row["sides"]) in sides: all_rows[name].append({"run":str(root),**row})
    wb=Workbook(); notes=wb.active; notes.title="metadata"; notes.append(["item","value"])
    for pair in(("Re",100),("D_eq",1),("area",math.pi/4),("circle_code",0),
                 ("shape_convention","Equal area; polygons have a horizontal top edge; body centered at origin"),
                 ("force_normalization","C=F/(0.5 rho U_ref^2 D_eq); stress reconstruction estimates"),
                 ("corner_limitation","Finite-thickness penalty geometry and fluid-side extrapolation do not resolve singular vertex stresses exactly"),
                 ("probe","Signed v/U_ref at x/D_eq=2, y/D_eq=0; each accepted time step"),
                 ("amplitude_window",f"{AMPLITUDE_START} to {AMPLITUDE_END}"),
                 ("amplitude_note","Half peak-to-peak in selected window; not proof of periodic saturation")):
        notes.append(pair)
    for name,title in(("history.csv","penalty_history"),("force_history.csv","surface_force_history"),("probe_history.csv","probe_history")):
        rows=all_rows[name]
        if not rows: continue
        sheet=wb.create_sheet(title); keys=list(rows[0]); sheet.append(keys)
        for row in rows:
            values=[]
            for key in keys:
                value=row.get(key)
                if key not in("run","shape"):
                    try:
                        number=float(value); value=number if math.isfinite(number) else str(value)
                    except(TypeError,ValueError): pass
                values.append(value)
            sheet.append(values)
    sheet=wb.create_sheet("probe_amplitude"); sheet.append(["run","shape","sides","Re","window_start","window_end","samples","Av_U","status"])
    groups={}
    for row in all_rows["probe_history.csv"]: groups.setdefault((row["run"],int(row["sides"])),[]).append(row)
    plt=pyplot(); fe,ae=plt.subplots(figsize=(10,5)); ff,af=plt.subplots(figsize=(9,5)); labels=[]; values=[]
    for (run,n),rows in groups.items():
        rows.sort(key=lambda r:float(r["time"])); t=np.array([float(r["time"]) for r in rows]); v=np.array([float(r["vp_U"]) for r in rows])
        if not np.isfinite(v).all() or np.any(np.diff(t)<=0): raise ValueError("Invalid probe history")
        ae.plot(t,v,label=shape_name(n)); mask=(t>=AMPLITUDE_START-1e-8)&(t<=AMPLITUDE_END+1e-8); tw=t[mask]; vw=v[mask]
        config_path=Path(run)/"config_used.json"; config=json.loads(config_path.read_text(encoding="utf-8")) if config_path.exists() else {}
        dt=config.get("physics",{}).get("dt")
        complete=(len(tw)>1 and math.isclose(tw[0],AMPLITUDE_START,abs_tol=1e-7) and math.isclose(tw[-1],AMPLITUDE_END,abs_tol=1e-7)
                  and dt is not None and np.allclose(np.diff(tw),dt,rtol=1e-4,atol=1e-7))
        amplitude=float(np.ptp(vw)/2) if complete else None
        sheet.append([run,shape_name(n),n,100,AMPLITUDE_START,AMPLITUDE_END,len(tw),amplitude,"window_complete_not_stability_certified" if complete else "incomplete_window"])
        if complete: labels.append(shape_name(n)); values.append(amplitude)
    ae.set(xlabel="t*=t U_ref/D_eq",ylabel="v_p/U_ref",title="(e) Probe transverse velocity, Re=100")
    if groups: ae.legend(fontsize=8)
    ae.axhline(0,color="gray",lw=.5)
    if values:
        af.scatter(range(len(values)),values); af.set_xticks(range(len(labels))); af.set_xticklabels(labels,rotation=30,ha="right")
    else: af.text(.5,.5,"No complete late-time window",transform=af.transAxes,ha="center")
    af.set(xlabel="Geometry (equal area)",ylabel="A_v/U_ref",title="(f) Late-window probe amplitude, Re=100"); af.set_ylim(bottom=0)
    for fig,name in((fe,"e_probe_velocity"),(ff,"f_probe_amplitude")):
        fig.tight_layout(); fig.savefig(destination/(name+".png"),dpi=180); plt.close(fig)
    for sh in wb:
        sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="24476B")
        for i,column in enumerate(sh.columns,1): sh.column_dimensions[get_column_letter(i)].width=min(65,max(14,max(len(str(c.value or "")) for c in column)+2))
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float): cell.number_format="0.00000000"
    wb.save(destination/"geometry_force_probe_data.xlsx")
    print("Excel and geometry comparisons:",destination,flush=True)


def plot_training_results(directory=None,options=None):
    root=settings.resolve_path(directory) if directory is not None else settings.resolve_path(TRAIN_RESULT_DIRECTORY) if TRAIN_RESULT_DIRECTORY else settings.latest_training_directory()
    rows=read_rows(root/"history.csv")
    if not rows: print("No accepted time steps are available for plotting:",root); return False
    if "sides" not in rows[0]: raise ValueError("This plotter requires geometry-version results")
    opts={"sides":PLOT_SIDES,"times":PLOT_TIMES,"discard_time":PLOT_DISCARD_TIME,"sequence_start":VORTICITY_START,
          "sequence_end":VORTICITY_END,"sequence_interval":VORTICITY_INTERVAL,"fps":VORTICITY_FPS,"gif":MAKE_VORTICITY_GIF}
    if options: opts.update(options)
    seq=requested_times(opts["sequence_start"],opts["sequence_end"],opts["sequence_interval"])
    if not opts["times"] or any(not math.isfinite(t) or t<0 for t in opts["times"]): raise ValueError("Invalid PLOT_TIMES")
    output=settings.resolve_path(FIGURE_DIRECTORY) if FIGURE_DIRECTORY and options is None else root/"plots_existing"
    output=output/datetime.now().strftime("%Y%m%d_%H%M%S_%f"); output.mkdir(parents=True)
    cases={int(r["case"]):int(r["sides"]) for r in rows}; reports=[]; complete=True
    for case,n in cases.items():
        if opts["sides"] is not None and n not in opts["sides"]: continue
        destination=output/shape_name(n); destination.mkdir(); fields=inventory(root,case,n)
        force_plots(root,case,n,destination,opts["discard_time"])
        selected=plot_vorticity_sequence(fields,case,n,destination/"selected_snapshots",opts["times"],make_gif=False,limit=VORTICITY_LIMIT,dpi=VORTICITY_DPI)
        sequence=plot_vorticity_sequence(fields,case,n,destination/"vorticity_sequence",seq,fps=opts["fps"],make_gif=opts["gif"],limit=VORTICITY_LIMIT,dpi=VORTICITY_DPI,interval=opts["sequence_interval"])
        cycle=phase_plot(root,case,n,fields,destination)
        complete=complete and not selected["missing_times"] and not sequence["missing_times"]
        reports.append({"sides":n,"selected":selected,"sequence":sequence,"cycle":cycle})
    if opts["sides"] is not None:
        missing=set(opts["sides"])-set(cases.values())
        if missing: complete=False; reports.append({"missing_shapes":sorted(missing)})
    export_comparisons([root],output,opts["discard_time"],sides=opts["sides"])
    (output/"plot_report.json").write_text(json.dumps(reports,indent=2),encoding="utf-8")
    print("Geometry-results plotting directory:",output,flush=True)
    return complete


if __name__=="__main__":
    raise SystemExit(0 if plot_training_results() else 2)
